```python
import io
from datetime import date, datetime

import pandas as pd
import streamlit as st

from openpyxl.styles import Alignment, Font, PatternFill

from config.schema import MASTER_SCHEMA
from utils.database import DatabaseManager


EXCEL_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)

EXCEL_FILENAME = "credential_records.xlsx"


DATABASE_TO_EXCEL_FIELDS = {
    "Sr. No.": None,
    "Client (PMA)": "client_pma",
    "System (LMBS, PMBS, MMBS, OLMRTS)": "system",
    "Contract": "contract",
    "Document Type": "document_type",
    "Document Number": "document_number",
    "Date of Issuance": "date_of_issuance",
    "Amount": "amount",
    "Initiated By": "initiated_by",
    "Reviewed By": "reviewed_by",
    "Approved by": "approved_by",
}


def _validate_database_record(record):
    """Validate the basic structure of a database record."""

    if not isinstance(record, dict):
        raise ValueError(
            "A database record is not in a valid format."
        )

    required_database_fields = {
        field
        for field in DATABASE_TO_EXCEL_FIELDS.values()
        if field is not None
    }

    missing_fields = (
        required_database_fields
        - set(record.keys())
    )

    if missing_fields:
        raise ValueError(
            "A database record is missing expected fields: "
            f"{sorted(missing_fields)}"
        )


def _normalize_excel_date(value):
    """Convert supported database date values to a Python date."""

    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    try:
        parsed = pd.to_datetime(
            value,
            errors="raise",
        )
        return parsed.date()

    except Exception as exc:
        raise ValueError(
            f"Invalid date value encountered: {value}"
        ) from exc


def _records_to_dataframe(records):
    """Convert database records to the exact locked master schema."""

    if records is None:
        records = []

    if not isinstance(records, list):
        raise ValueError(
            "Loaded credential records must be a list."
        )

    rows = []

    for index, record in enumerate(
        records,
        start=1,
    ):
        _validate_database_record(
            record
        )

        row = {
            "Sr. No.": index,
        }

        for excel_field, database_field in (
            DATABASE_TO_EXCEL_FIELDS.items()
        ):
            if database_field is None:
                continue

            value = record.get(
                database_field
            )

            if excel_field == "Date of Issuance":
                value = _normalize_excel_date(
                    value
                )

            row[excel_field] = value

        rows.append(row)

    dataframe = pd.DataFrame(
        rows,
        columns=MASTER_SCHEMA,
    )

    if list(dataframe.columns) != MASTER_SCHEMA:
        raise ValueError(
            "Generated records dataframe does not match "
            "the locked master schema."
        )

    return dataframe


def _create_excel_file(dataframe):
    """Create a formatted Excel workbook in memory."""

    if not isinstance(
        dataframe,
        pd.DataFrame,
    ):
        raise TypeError(
            "Excel export requires a pandas DataFrame."
        )

    if list(dataframe.columns) != MASTER_SCHEMA:
        raise ValueError(
            "Excel export columns do not match "
            "the locked master schema."
        )

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        dataframe.to_excel(
            writer,
            index=False,
            sheet_name="Credentials",
        )

        worksheet = writer.sheets[
            "Credentials"
        ]

        # ----------------------------------------------------
        # Worksheet usability
        # ----------------------------------------------------

        worksheet.freeze_panes = "A2"

        if worksheet.max_row >= 1:
            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

        # ----------------------------------------------------
        # Header formatting
        # ----------------------------------------------------

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        header_font = Font(
            bold=True,
            color="FFFFFF",
        )

        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        worksheet.row_dimensions[1].height = 30

        # ----------------------------------------------------
        # Column widths
        # ----------------------------------------------------

        column_widths = {
            "A": 10,
            "B": 25,
            "C": 28,
            "D": 22,
            "E": 24,
            "F": 22,
            "G": 18,
            "H": 18,
            "I": 25,
            "J": 25,
            "K": 25,
        }

        for column, width in column_widths.items():
            worksheet.column_dimensions[
                column
            ].width = width

        # ----------------------------------------------------
        # Date formatting
        # ----------------------------------------------------

        for cell in worksheet["G"][1:]:
            cell.number_format = "yyyy-mm-dd"

        # ----------------------------------------------------
        # Amount formatting
        # ----------------------------------------------------

        for cell in worksheet["H"][1:]:
            cell.number_format = "#,##0.00"

        # ----------------------------------------------------
        # Alignment
        # ----------------------------------------------------

        for cell in worksheet["A"][1:]:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for cell in worksheet["G"][1:]:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for cell in worksheet["H"][1:]:
            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
            )

        # Wrap longer text fields.
        for row in worksheet.iter_rows(
            min_row=2,
        ):
            for cell in row:
                if cell.column not in {
                    1,
                    7,
                    8,
                }:
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )

    output.seek(0)

    return output.getvalue()


def load_user_records():
    """Load records belonging to the currently authenticated user."""

    access_token = st.session_state.get(
        "access_token"
    )

    refresh_token = st.session_state.get(
        "refresh_token"
    )

    if not access_token:
        raise ValueError(
            "Your authentication session has expired. "
            "Please log in again."
        )

    if not refresh_token:
        raise ValueError(
            "Your authentication session is incomplete. "
            "Please log in again."
        )

    database = DatabaseManager()

    database.set_user_session(
        access_token,
        refresh_token,
    )

    return database.get_user_credentials()


def render_records_ui():
    """Render the user's saved credential records and Excel export."""

    st.divider()

    st.header(
        "📋 My Credential Records"
    )

    st.write(
        "These are the credential records permanently "
        "saved in your Supabase database."
    )

    # --------------------------------------------------------
    # Load records
    # --------------------------------------------------------

    if st.button(
        "🔄 Load My Records",
        key="load_my_records_button",
    ):

        try:
            with st.spinner(
                "Loading your records..."
            ):
                records = load_user_records()

            # Always replace the previously loaded data.
            st.session_state.my_credentials = (
                records
            )

            if not records:
                st.info(
                    "No saved credential records found."
                )
            else:
                st.success(
                    f"{len(records)} credential "
                    "record(s) loaded successfully."
                )

        except ValueError as exc:

            # Clear potentially stale records if the
            # current authentication/session is invalid.
            st.session_state.my_credentials = []

            st.error(
                f"Unable to load records: {exc}"
            )

        except Exception:

            st.session_state.my_credentials = []

            st.error(
                "Your credential records could not be loaded. "
                "Please try again."
            )

    records = st.session_state.get(
        "my_credentials"
    )

    if not records:
        return

    # --------------------------------------------------------
    # Build dataframe
    # --------------------------------------------------------

    try:
        dataframe = _records_to_dataframe(
            records
        )

    except ValueError as exc:

        st.error(
            f"Unable to prepare your records: {exc}"
        )

        return

    except Exception:

        st.error(
            "Unable to prepare the loaded records."
        )

        return

    # --------------------------------------------------------
    # Display records
    # --------------------------------------------------------

    st.dataframe(
        dataframe,
        use_container_width=True,
        hide_index=True,
    )

    # --------------------------------------------------------
    # Excel export
    # --------------------------------------------------------

    st.subheader(
        "📥 Export Records"
    )

    st.write(
        "Download your currently loaded credential "
        "records as an Excel workbook."
    )

    try:
        excel_file = _create_excel_file(
            dataframe
        )

        st.download_button(
            label="📊 Download Excel",
            data=excel_file,
            file_name=EXCEL_FILENAME,
            mime=EXCEL_MIME_TYPE,
            type="primary",
            key="download_credentials_excel",
        )

    except Exception:

        st.error(
            "Unable to create the Excel export. "
            "Please try loading the records again."
        )
```
